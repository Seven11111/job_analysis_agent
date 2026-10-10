"""Safe, bounded text extraction for locally uploaded resume documents."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


MAX_RESUME_BYTES = 10 * 1024 * 1024
MAX_RESUME_TEXT_CHARS = 80_000
MAX_PDF_PAGES = 60
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 600


class ResumeDocumentError(ValueError):
    """The uploaded file is unsupported or cannot be safely read as a resume."""


def extract_resume_text(filename: str, content: bytes) -> tuple[str, str, str]:
    """Return extracted text, safe extension, and normalized content type."""

    suffix = Path(filename.replace("\\", "/")).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise ResumeDocumentError("目前只支持 PDF 和 DOCX 格式的简历。")
    if not content:
        raise ResumeDocumentError("上传文件为空，请选择有效的简历文件。")
    if len(content) > MAX_RESUME_BYTES:
        raise ResumeDocumentError("简历文件不能超过 10 MB。")

    if suffix == ".pdf":
        text = _extract_pdf(content)
        content_type = "application/pdf"
    else:
        text = _extract_docx(content)
        content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
    if not text:
        raise ResumeDocumentError("未能从文件中提取到文字；扫描版 PDF 暂不支持，请上传可复制文字的 PDF 或 DOCX。")
    if len(text) > MAX_RESUME_TEXT_CHARS:
        raise ResumeDocumentError("简历文字内容超过解析上限，请精简文件后重试。")
    return text, suffix, content_type


def _extract_pdf(content: bytes) -> str:
    if not content.startswith(b"%PDF-"):
        raise ResumeDocumentError("文件内容不是有效的 PDF。")
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ResumeDocumentError("暂不支持加密 PDF，请先解除密码保护后重试。")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ResumeDocumentError("PDF 页数不能超过 60 页。")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ResumeDocumentError:
        raise
    except Exception as exc:
        raise ResumeDocumentError("PDF 解析失败，请确认文件未损坏且可以复制文字。") from exc


def _extract_docx(content: bytes) -> str:
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_ENTRIES or sum(item.file_size for item in infos) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ResumeDocumentError("DOCX 文件内容异常或过大，已停止解析。")
            if "word/document.xml" not in archive.namelist():
                raise ResumeDocumentError("文件内容不是有效的 DOCX 文档。")
        from docx import Document

        document = Document(BytesIO(content))
        pieces = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                pieces.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(pieces)
    except ResumeDocumentError:
        raise
    except Exception as exc:
        raise ResumeDocumentError("DOCX 解析失败，请确认文件未损坏。") from exc
